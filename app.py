from __future__ import annotations

import io
import json
import re
import traceback
from copy import copy
from typing import Any, Dict, List, Tuple

import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

APP_VERSION = "V2.1 Recovery Stable"

st.set_page_config(page_title="Excel Formula Fusion", layout="wide")

# Keep CSS minimal. Aggressive CSS caused invisible upload/select text in earlier builds.
st.markdown(
    """
    <style>
    .stButton button, .stDownloadButton button {
        background-color: #f36f21 !important;
        color: white !important;
        font-weight: 700 !important;
        border: 1px solid #c95512 !important;
    }
    .stButton button *, .stDownloadButton button * { color: white !important; }
    </style>
    """,
    unsafe_allow_html=True,
)

# ---------------- helpers ----------------

def norm(v: Any) -> str:
    return "" if v is None else str(v).strip()


def safe_col(v: str, default: str) -> str:
    v = (v or "").strip().upper()
    if re.fullmatch(r"[A-Z]{1,3}", v):
        return v
    return default


def sheet_default(sheets: List[str], wanted: str, fallback: int = 0) -> str:
    for s in sheets:
        if s.strip().lower() == wanted.strip().lower():
            return s
    return sheets[min(fallback, max(0, len(sheets) - 1))]


def quote_sheet(name: str) -> str:
    return name.replace("'", "''")


def apply_heading(cell):
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.font = Font(color="FFFFFF", bold=True)


def copy_row_style(ws, source_row: int, target_row: int, start_col: int, end_col: int) -> None:
    for c in range(start_col, end_col + 1):
        src = ws.cell(source_row, c)
        dst = ws.cell(target_row, c)
        if src.has_style:
            dst.font = copy(src.font)
            dst.fill = copy(src.fill)
            dst.border = copy(src.border)
            dst.alignment = copy(src.alignment)
            dst.number_format = src.number_format
            dst.protection = copy(src.protection)


def parse_size_to_sqm(size_text: str) -> float:
    s = norm(size_text).lower().replace("×", "x").replace(",", "")
    nums = re.findall(r"(\d+(?:\.\d+)?)", s)
    if len(nums) >= 2:
        w, h = float(nums[0]), float(nums[1])
        if "cm" in s and "mm" not in s:
            return (w / 100.0) * (h / 100.0)
        if re.search(r"\bm\b", s) and "mm" not in s and "cm" not in s:
            return w * h
        return (w / 1000.0) * (h / 1000.0)
    if len(nums) == 1 and ("dia" in s or "ø" in s or "round" in s):
        d = float(nums[0])
        if "cm" in s and "mm" not in s:
            d_m = d / 100.0
        elif re.search(r"\bm\b", s) and "mm" not in s and "cm" not in s:
            d_m = d
        else:
            d_m = d / 1000.0
        return 3.141592653589793 * (d_m / 2.0) ** 2
    return 0.0


def detect_multiplier(name: str) -> Tuple[int, str, str]:
    t = norm(name).upper()
    m = re.search(r"\bSET\s*(?:OF)?\s*(\d{1,3})\b", t)
    if m:
        n = int(m.group(1))
        if 1 < n <= 500:
            return n, "red", f"set of {n}"
    m = re.search(r"\b(?:1\s*)?PACK\s*=\s*(\d{1,5})\b", t)
    if m:
        n = int(m.group(1))
        if 1 < n <= 10000:
            return n, "red", f"pack = {n}"
    m = re.search(r"\bPACK\s+OF\s+(\d{1,5})\b", t)
    if m:
        n = int(m.group(1))
        if 1 < n <= 10000:
            return n, "red", f"pack of {n}"
    if "SET" in t or "PACK" in t:
        return 1, "orange", "set/pack wording unclear"
    return 1, "none", ""


def ignored_array(countries: List[str]) -> str:
    vals = [c.strip().upper().replace('"', "") for c in countries if c.strip()]
    if not vals:
        vals = ["NZ"]
    return "{" + ",".join(f'"{v}"' for v in vals) + "}"


def build_clean_qty_formula(col: str, total_row: int, scan_start: int, scan_end: int, country_col: str, ignored: List[str], multiplier: int) -> str:
    arr = ignored_array(ignored)
    # total qty - qty where country is ignored; do not include blank country rows.
    base = (
        f"({col}${total_row}-SUMPRODUCT(({col}${scan_start}:{col}${scan_end})*"
        f"(--ISNUMBER(MATCH(UPPER(${country_col}${scan_start}:${country_col}${scan_end}),{arr},0)))))"
    )
    if multiplier != 1:
        return f"={base}*{multiplier}"
    return f"={base}"


def build_ds_formula(col: str, name_row: int, size_row: int, ref_sheet: str, ref_name_col: str, ref_size_col: str, ref_ds_col: str, ref_start: int, ref_end: int) -> str:
    rs = quote_sheet(ref_sheet)
    return (
        f"=IFERROR(INDEX('{rs}'!${ref_ds_col}${ref_start}:${ref_ds_col}${ref_end},"
        f"MATCH(1,INDEX(('{rs}'!${ref_name_col}${ref_start}:${ref_name_col}${ref_end}={col}${name_row})*"
        f"('{rs}'!${ref_size_col}${ref_start}:${ref_size_col}${ref_end}={col}${size_row}),0),0)),\"\")"
    )


def get_stock_options(ws_values, stock_row: int, start_col: int, end_col: int) -> List[str]:
    out, seen = [], set()
    for c in range(start_col, end_col + 1):
        s = norm(ws_values.cell(stock_row, c).value)
        if not s:
            continue
        key = s.upper()
        if key not in seen:
            seen.add(key)
            out.append(s)
    return out


def detect_country_end(ws_values, country_col: str, start_row: int, limit: int = 400) -> int:
    tokens = {"NZ", "AUS", "AU", "AUSTRALIA", "NEW ZEALAND", "FIJI", "SG", "SINGAPORE"}
    cidx = column_index_from_string(country_col)
    last = start_row
    for r in range(start_row, min(ws_values.max_row, limit) + 1):
        if norm(ws_values.cell(r, cidx).value).upper() in tokens:
            last = r
    return max(last, start_row)


def build_workbook(upload_bytes: bytes, cfg: Dict[str, Any], rates: Dict[str, float]) -> bytes:
    wb = load_workbook(io.BytesIO(upload_bytes))
    wbv = load_workbook(io.BytesIO(upload_bytes), data_only=True, read_only=True)
    ws = wb[cfg["working_sheet"]]
    wsv = wbv[cfg["working_sheet"]]

    start_col = column_index_from_string(cfg["item_start_col"])
    end_col = column_index_from_string(cfg["item_end_col"])
    selected = set(cfg.get("selected_stocks", []))
    ds_factor = 1.0 + (float(cfg.get("ds_loading", 20.0)) / 100.0)

    red_fill = PatternFill("solid", fgColor="FFC7CE")
    orange_fill = PatternFill("solid", fgColor="FCE4D6")
    green_fill = PatternFill("solid", fgColor="C6EFCE")
    yellow_fill = PatternFill("solid", fgColor="FFF2CC")

    labels = [
        (cfg["ds_row"], "DS/SS Lookup"),
        (cfg["clean_qty_row"], "Clean Qty"),
        (cfg["multiplier_row"], "Qty Multiplier"),
        (cfg["sqm_row"], "SQM"),
        (cfg["price_row"], "Price"),
    ]
    for row, label in labels:
        ws.cell(row, 1).value = label
        apply_heading(ws.cell(row, 1))
    copy_row_style(ws, cfg["total_qty_row"], cfg["clean_qty_row"], start_col, end_col)

    audit = []
    for c in range(start_col, end_col + 1):
        col = get_column_letter(c)
        name = norm(wsv.cell(cfg["name_row"], c).value or ws.cell(cfg["name_row"], c).value)
        size = norm(wsv.cell(cfg["size_row"], c).value or ws.cell(cfg["size_row"], c).value)
        stock = norm(wsv.cell(cfg["stock_row"], c).value or ws.cell(cfg["stock_row"], c).value)
        mult, flag, reason = detect_multiplier(name)

        ws.cell(cfg["ds_row"], c).value = build_ds_formula(
            col, cfg["name_row"], cfg["size_row"], cfg["reference_sheet"], cfg["ref_name_col"], cfg["ref_size_col"], cfg["ref_ds_col"], cfg["ref_start_row"], cfg["ref_end_row"]
        )
        ws.cell(cfg["clean_qty_row"], c).value = build_clean_qty_formula(
            col, cfg["total_qty_row"], cfg["country_scan_start"], cfg["country_scan_end"], cfg["country_col"], cfg["ignored_countries"], mult
        )
        ws.cell(cfg["multiplier_row"], c).value = mult

        sqm_each = parse_size_to_sqm(size)
        ws.cell(cfg["sqm_row"], c).value = f"={col}${cfg['clean_qty_row']}*{sqm_each:.6f}" if sqm_each else ""

        rate = float(rates.get(stock, 0.0))
        if stock in selected and rate > 0:
            ws.cell(cfg["price_row"], c).value = f'=IF(UPPER({col}${cfg["ds_row"]})="DS",{col}${cfg["sqm_row"]}*{rate}*{ds_factor},{col}${cfg["sqm_row"]}*{rate})'
        else:
            ws.cell(cfg["price_row"], c).value = ""

        if flag == "red":
            ws.cell(cfg["name_row"], c).fill = red_fill
            ws.cell(cfg["clean_qty_row"], c).fill = red_fill
            ws.cell(cfg["multiplier_row"], c).fill = red_fill
            audit.append([col, name, size, stock, mult, "RED", reason])
        elif flag == "orange":
            ws.cell(cfg["name_row"], c).fill = orange_fill
            ws.cell(cfg["multiplier_row"], c).fill = orange_fill
            audit.append([col, name, size, stock, mult, "ORANGE", reason])
        else:
            ws.cell(cfg["multiplier_row"], c).fill = green_fill
        if stock in selected:
            ws.cell(cfg["stock_row"], c).fill = yellow_fill

    if "Qty Multiplier Audit" in wb.sheetnames:
        del wb["Qty Multiplier Audit"]
    aud = wb.create_sheet("Qty Multiplier Audit")
    aud.append(["Column", "Name", "Size", "Stock", "Multiplier", "Flag", "Reason"])
    for cell in aud[1]:
        apply_heading(cell)
    for row in audit:
        aud.append(row)
        fill = red_fill if row[5] == "RED" else orange_fill
        for cell in aud[aud.max_row]:
            cell.fill = fill

    if "Stock SQM Summary" in wb.sheetnames:
        del wb["Stock SQM Summary"]
    sm = wb.create_sheet("Stock SQM Summary")
    sm.append(["Stock", "Rate / sqm", "Total SQM", "Estimated Price", "DS Loading %"])
    for cell in sm[1]:
        apply_heading(cell)
    work = quote_sheet(cfg["working_sheet"])
    r = 2
    for stock in cfg.get("selected_stocks", []):
        sm.cell(r, 1).value = stock
        sm.cell(r, 2).value = float(rates.get(stock, 0.0))
        sm.cell(r, 3).value = f'=SUMIF(\'{work}\'!${cfg["item_start_col"]}${cfg["stock_row"]}:${cfg["item_end_col"]}${cfg["stock_row"]},A{r},\'{work}\'!${cfg["item_start_col"]}${cfg["sqm_row"]}:${cfg["item_end_col"]}${cfg["sqm_row"]})'
        sm.cell(r, 4).value = f'=SUMIF(\'{work}\'!${cfg["item_start_col"]}${cfg["stock_row"]}:${cfg["item_end_col"]}${cfg["stock_row"]},A{r},\'{work}\'!${cfg["item_start_col"]}${cfg["price_row"]}:${cfg["item_end_col"]}${cfg["price_row"]})'
        sm.cell(r, 5).value = cfg.get("ds_loading", 20.0)
        r += 1

    for sheet in [aud, sm]:
        for col_idx in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col_idx)].width = 26

    try:
        wbv.close()
    except Exception:
        pass
    out = io.BytesIO()
    wb.save(out)
    try:
        wb.close()
    except Exception:
        pass
    return out.getvalue()


# ---------------- UI ----------------

st.title("Excel Formula Fusion")
st.caption(APP_VERSION)

for key, default in {
    "uploaded_bytes": None,
    "uploaded_name": "",
    "sheet_names": [],
    "rate_memory": {},
}.items():
    if key not in st.session_state:
        st.session_state[key] = default

with st.sidebar:
    st.header("Upload")
    uploaded = st.file_uploader("Upload Excel workbook", type=["xlsx", "xlsm"])
    if uploaded is not None:
        st.session_state.uploaded_bytes = uploaded.getvalue()
        st.session_state.uploaded_name = uploaded.name
        st.success(f"Selected: {uploaded.name}")

    rates_file = st.file_uploader("Optional stock rate JSON", type=["json"])
    if rates_file is not None:
        try:
            st.session_state.rate_memory.update(json.loads(rates_file.getvalue().decode("utf-8")))
            st.success("Rate memory loaded")
        except Exception as exc:
            st.error(f"Could not read rate JSON: {exc}")

if not st.session_state.uploaded_bytes:
    st.info("Upload an Excel file. Nothing heavy runs until you press the read/generate buttons.")
    st.stop()

st.success(f"Workbook uploaded: {st.session_state.uploaded_name}")
st.write(f"Size: {len(st.session_state.uploaded_bytes) / 1024 / 1024:.2f} MB")

if st.button("Read workbook / detect sheets"):
    try:
        with st.spinner("Reading sheet names only..."):
            wb = load_workbook(io.BytesIO(st.session_state.uploaded_bytes), read_only=True, data_only=True)
            st.session_state.sheet_names = wb.sheetnames
            wb.close()
        st.success("Workbook read successfully.")
    except Exception:
        st.error("Failed while reading workbook sheet names.")
        st.code(traceback.format_exc())

if not st.session_state.sheet_names:
    st.warning("Press **Read workbook / detect sheets** to continue.")
    st.stop()

sheets = st.session_state.sheet_names
st.write("Sheets detected:", ", ".join(sheets))

wd = sheet_default(sheets, "DL ANZ ALLOCATION", 0)
rd = sheet_default(sheets, "PRINT DB", 1 if len(sheets) > 1 else 0)

with st.form("mapping_form"):
    st.subheader("Mapping")
    c1, c2 = st.columns(2)
    with c1:
        working_sheet = st.selectbox("Working sheet", sheets, index=sheets.index(wd))
        name_row = int(st.number_input("Name row", min_value=1, value=4))
        size_row = int(st.number_input("Size row", min_value=1, value=5))
        stock_row = int(st.number_input("Stock/material row", min_value=1, value=6))
        total_qty_row = int(st.number_input("Original total qty row", min_value=1, value=7))
        country_col = safe_col(st.text_input("Country column", value="I"), "I")
        ignored = st.text_input("Ignore countries", value="NZ")
    with c2:
        reference_sheet = st.selectbox("Reference sheet", sheets, index=sheets.index(rd))
        ref_name_col = safe_col(st.text_input("Reference name column", value="C"), "C")
        ref_size_col = safe_col(st.text_input("Reference size column", value="E"), "E")
        ref_ds_col = safe_col(st.text_input("Reference DS/SS column", value="F"), "F")
        ref_start_row = int(st.number_input("Reference start row", min_value=1, value=12))
        ref_end_row = int(st.number_input("Reference end row", min_value=1, value=141))

    st.subheader("Item columns and output rows")
    m1, m2, m3, m4 = st.columns(4)
    with m1:
        item_start_col = safe_col(st.text_input("Start column", value="AC"), "AC")
        item_end_col = safe_col(st.text_input("End column", value="IG"), "IG")
    with m2:
        ds_row = int(st.number_input("DS/SS output row", min_value=1, value=168))
        clean_qty_row = int(st.number_input("Clean qty row", min_value=1, value=169))
    with m3:
        multiplier_row = int(st.number_input("Multiplier row", min_value=1, value=170))
        sqm_row = int(st.number_input("SQM row", min_value=1, value=171))
    with m4:
        price_row = int(st.number_input("Price row", min_value=1, value=172))
        ds_loading = float(st.number_input("DS loading %", min_value=0.0, value=20.0, step=1.0))

    default_rate = float(st.number_input("Default new stock rate", min_value=0.0, value=0.0, step=0.1, format="%.2f"))
    applied = st.form_submit_button("Apply mapping / refresh stock list")

if applied or "cfg" not in st.session_state:
    try:
        start_i = column_index_from_string(item_start_col)
        end_i = column_index_from_string(item_end_col)
        if end_i < start_i:
            raise ValueError("End column must be after start column.")
        wbv = load_workbook(io.BytesIO(st.session_state.uploaded_bytes), read_only=True, data_only=True)
        wsv = wbv[working_sheet]
        stock_options = get_stock_options(wsv, stock_row, start_i, end_i)
        scan_start = total_qty_row + 1
        scan_end = detect_country_end(wsv, country_col, scan_start)
        wbv.close()
        st.session_state.cfg = {
            "working_sheet": working_sheet,
            "reference_sheet": reference_sheet,
            "name_row": name_row,
            "size_row": size_row,
            "stock_row": stock_row,
            "total_qty_row": total_qty_row,
            "country_col": country_col,
            "ignored_countries": [x.strip().upper() for x in ignored.split(",") if x.strip()],
            "ref_name_col": ref_name_col,
            "ref_size_col": ref_size_col,
            "ref_ds_col": ref_ds_col,
            "ref_start_row": ref_start_row,
            "ref_end_row": ref_end_row,
            "item_start_col": item_start_col,
            "item_end_col": item_end_col,
            "ds_row": ds_row,
            "clean_qty_row": clean_qty_row,
            "multiplier_row": multiplier_row,
            "sqm_row": sqm_row,
            "price_row": price_row,
            "ds_loading": ds_loading,
            "country_scan_start": scan_start,
            "country_scan_end": scan_end,
            "stock_options": stock_options,
            "selected_stocks": st.session_state.get("selected_stocks", []),
        }
        st.success(f"Mapping applied. Stock scan: {item_start_col}:{item_end_col}. Country qty scan rows: {scan_start}:{scan_end}. Stocks found: {len(stock_options)}")
    except Exception:
        st.error("Mapping failed.")
        st.code(traceback.format_exc())

if "cfg" not in st.session_state:
    st.stop()

cfg = st.session_state.cfg
with st.expander("Current mapping", expanded=False):
    show_cfg = {k: v for k, v in cfg.items() if k != "stock_options"}
    st.write(show_cfg)

st.subheader("Stock/material rates")
st.caption(f"Stock names are read from row {cfg['stock_row']} between {cfg['item_start_col']} and {cfg['item_end_col']} only.")
selected_stocks = st.multiselect("Pick stock/material to calculate SQM and rate", cfg.get("stock_options", []), default=cfg.get("selected_stocks", []))
st.session_state.selected_stocks = selected_stocks
st.session_state.cfg["selected_stocks"] = selected_stocks

if selected_stocks:
    with st.form("rates_form"):
        staged = {}
        cols = st.columns(2)
        for i, stock in enumerate(selected_stocks):
            with cols[i % 2]:
                current = float(st.session_state.rate_memory.get(stock, default_rate))
                staged[stock] = st.number_input(f"Rate $/sqm — {stock}", min_value=0.0, value=current, step=0.1, format="%.2f", key=f"rate_{i}_{abs(hash(stock))}")
        save_rates = st.form_submit_button("Refresh / Update Rates")
        if save_rates:
            st.session_state.rate_memory.update(staged)
            st.success("Rates saved for this session.")
else:
    st.info("Select stock/material names above to show rate entry fields.")

st.download_button(
    "Download stock rate memory JSON",
    data=json.dumps(st.session_state.rate_memory, indent=2).encode("utf-8"),
    file_name="stock_rate_memory.json",
    mime="application/json",
)

st.subheader("Generate")
st.warning("Only press Generate after mapping and rates are correct. If Streamlit Cloud still runs out of memory, narrow the Start/End column range.")
if st.button("Generate Excel Workbook"):
    try:
        with st.spinner("Generating workbook..."):
            out = build_workbook(st.session_state.uploaded_bytes, st.session_state.cfg, st.session_state.rate_memory)
        st.success("Workbook generated.")
        st.download_button(
            "Download Excel Workbook",
            data=out,
            file_name=f"formula_fusion_{st.session_state.uploaded_name}",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    except Exception:
        st.error("Generation failed.")
        st.code(traceback.format_exc())
